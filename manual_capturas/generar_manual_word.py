# -*- coding: utf-8 -*-
"""
Generador Maestro del Documento Word Oficial:
DIRECCIÓN TIC - MANUAL TÉCNICO DE ARQUITECTURA FRONTEND Y MANUAL DE USUARIO
APLICATIVO CARNET DIGITAL (VERSIÓN 2026)
Estandarizado según la Guía de Estilo de Escritura de Microsoft (Microsoft Writing Style Guide)
"""

import os
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

DOCX_OUTPUT_PATH = r"f:\personal\CoreCredential-Front-main\MANUAL_TECNICO_Y_DE_USUARIO_CARNET_DIGITAL.docx"
IMAGES_DIR = r"f:\personal\CoreCredential-Front-main\manual_capturas\recortadas"

def set_cell_background(cell, hex_color):
    """Establece el color de fondo de una celda de tabla."""
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)

def set_cell_margins(cell, top=120, bottom=120, left=160, right=160):
    """Establece el relleno interno de una celda en dxa."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)

def add_callout(doc, text, callout_type="TIP"):
    """
    Inserta un cuadro destacado estilizado.
    Sigue las pautas de Microsoft: voz activa, conciso y visualmente diferenciado.
    """
    bg_colors = {
        "TIP": "F0FDF4",      # Verde suave
        "NOTE": "F8FAFC",     # Azul/gris suave
        "WARNING": "FFFBEB",  # Ámbar suave
        "ALERT": "FEF2F2"     # Rojo suave
    }
    border_colors = {
        "TIP": "16A34A",
        "NOTE": "2563EB",
        "WARNING": "D97706",
        "ALERT": "DC2626"
    }
    prefixes = {
        "TIP": "Sugerencia: ",
        "NOTE": "Nota técnica: ",
        "WARNING": "Advertencia: ",
        "ALERT": "Acción crítica: "
    }

    bg = bg_colors.get(callout_type, "F8FAFC")
    bcolor = border_colors.get(callout_type, "2563EB")
    prefix = prefixes.get(callout_type, "Nota: ")

    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    cell.width = Inches(6.5)
    set_cell_background(cell, bg)
    set_cell_margins(cell, top=140, bottom=140, left=200, right=200)

    # Borde izquierdo destacado
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'<w:top w:val="none"/>'
        f'<w:left w:val="single" w:sz="36" w:space="0" w:color="{bcolor}"/>'
        f'<w:bottom w:val="none"/>'
        f'<w:right w:val="none"/>'
        f'</w:tcBorders>'
    )
    tcPr.append(tcBorders)

    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    
    r_prefix = p.add_run(prefix)
    r_prefix.bold = True
    r_prefix.font.size = Pt(9.5)
    r_prefix.font.color.rgb = RGBColor(0, 31, 63)
    
    r_text = p.add_run(text)
    r_text.font.size = Pt(9.5)
    r_text.font.color.rgb = RGBColor(51, 65, 85)

    doc.add_paragraph().paragraph_format.space_after = Pt(4)

def format_table(table, col_widths, headers, data, header_bg="0062E3"):
    """Aplica formato estructurado y legible a tablas institucionales."""
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    
    # Encabezados
    hdr_cells = table.rows[0].cells
    for i, title in enumerate(headers):
        hdr_cells[i].text = title
        set_cell_background(hdr_cells[i], header_bg)
        set_cell_margins(hdr_cells[i], top=120, bottom=120, left=140, right=140)
        p = hdr_cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        for run in p.runs:
            run.font.bold = True
            run.font.size = Pt(9)
            run.font.color.rgb = RGBColor(255, 255, 255)
            run.font.name = "Calibri"

    # Filas de datos
    for row_idx, row_data in enumerate(data):
        row = table.add_row()
        bg_row = "FFFFFF" if row_idx % 2 == 0 else "F8FAFC"
        for col_idx, val in enumerate(row_data):
            cell = row.cells[col_idx]
            cell.text = str(val)
            set_cell_background(cell, bg_row)
            set_cell_margins(cell, top=100, bottom=100, left=140, right=140)
            p = cell.paragraphs[0]
            for run in p.runs:
                run.font.size = Pt(8.5)
                run.font.color.rgb = RGBColor(51, 65, 85)
                run.font.name = "Calibri"

    # Anchos de columna
    for row in table.rows:
        for i, w in enumerate(col_widths):
            row.cells[i].width = Inches(w)

def add_screenshot(doc, img_name, fig_number, caption, max_width_inches=6.0):
    """
    Inserta la imagen con formato institucional según la guía de estilo:
    Figura X. [Descripción precisa en modo oración]
    """
    img_path = os.path.join(IMAGES_DIR, img_name)
    if not os.path.exists(img_path):
        p_err = doc.add_paragraph(f"[Imagen no encontrada: {img_name}]")
        p_err.runs[0].font.color.rgb = RGBColor(220, 38, 38)
        return

    p_img = doc.add_paragraph()
    p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_img.paragraph_format.space_before = Pt(8)
    p_img.paragraph_format.space_after = Pt(2)
    run_img = p_img.add_run()
    run_img.add_picture(img_path, width=Inches(max_width_inches))

    p_cap = doc.add_paragraph()
    p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_cap.paragraph_format.space_after = Pt(10)
    
    r_fig = p_cap.add_run(f"Figura {fig_number}. ")
    r_fig.bold = True
    r_fig.font.size = Pt(8.5)
    r_fig.font.color.rgb = RGBColor(0, 31, 63)
    
    r_cap = p_cap.add_run(caption)
    r_cap.font.size = Pt(8.5)
    r_cap.font.color.rgb = RGBColor(100, 116, 139)

def main():
    print("Creando documento Word institucional de Carnet Digital...")
    doc = Document()

    # Configuración de márgenes estándar (1 pulgada = 2.54 cm)
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    # =========================================================================
    # PORTADA INSTITUCIONAL
    # =========================================================================
    p_inst = doc.add_paragraph()
    p_inst.paragraph_format.space_before = Pt(36)
    p_inst.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_inst = p_inst.add_run("DIRECCIÓN TIC\nDEPARTAMENTO DE DESARROLLO E INNOVACIÓN TECNOLÓGICA")
    r_inst.bold = True
    r_inst.font.size = Pt(11)
    r_inst.font.color.rgb = RGBColor(0, 98, 227)

    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(24)
    p_title.paragraph_format.space_after = Pt(8)
    r_title = p_title.add_run("Carnet Digital")
    r_title.bold = True
    r_title.font.size = Pt(38)
    r_title.font.name = "Calibri"
    r_title.font.color.rgb = RGBColor(0, 31, 63)

    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.paragraph_format.space_after = Pt(24)
    r_sub = p_sub.add_run("Manual técnico de arquitectura frontend y guía oficial de uso")
    r_sub.font.size = Pt(16)
    r_sub.font.color.rgb = RGBColor(71, 85, 105)

    p_tag = doc.add_paragraph()
    p_tag.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_tag.paragraph_format.space_after = Pt(36)
    r_tag = p_tag.add_run("Plataforma empresarial de emisión, gestión y validación de credenciales inteligentes")
    r_tag.italic = True
    r_tag.font.size = Pt(10.5)
    r_tag.font.color.rgb = RGBColor(100, 116, 139)

    # Tabla de Control de Versiones Institucional (Estilo Modelo RH360)
    doc.add_heading("Control de versiones", level=2)
    tbl_version = doc.add_table(rows=1, cols=6)
    headers_version = ["Fecha", "Versión", "Autor / Elaboró", "Revisó", "Modificación", "Aprobación"]
    data_version = [
        (
            "28/09/2026",
            "2.5.0",
            "Ing. Especialista TIC\nArquitectura Frontend",
            "Dirección de Planeación y TIC",
            "Estandarización bajo Microsoft Writing Style Guide, transición de identidad a Carnet Digital y documentación técnica de 35 vistas.",
            "Dirección TIC /\nRectoría"
        )
    ]
    format_table(tbl_version, [0.8, 0.6, 1.3, 1.2, 1.6, 1.0], headers_version, data_version, "001F3F")

    doc.add_paragraph().paragraph_format.space_after = Pt(18)

    # Bloque formal de firmas y responsabilidades
    tbl_firmas = doc.add_table(rows=2, cols=3)
    tbl_firmas.alignment = WD_TABLE_ALIGNMENT.CENTER
    firmas_data = [
        ("Elaboró / Modificó:", "Revisó:", "Aprobó:"),
        (
            "Ing. Especialista TIC\nDesarrollo Web Frontend\nDirección TIC\nFecha: 28 de septiembre de 2026",
            "Directora de Planeación y TIC\nDirección de Planeación (E)\nFecha: 28 de septiembre de 2026",
            "Dirección General / Rectoría\nDirección Ejecutiva Institucional\nFecha: 28 de septiembre de 2026"
        )
    ]
    for row_idx, row_values in enumerate(firmas_data):
        row = tbl_firmas.rows[row_idx]
        for col_idx, text_val in enumerate(row_values):
            cell = row.cells[col_idx]
            cell.text = text_val
            cell.width = Inches(2.16)
            if row_idx == 0:
                set_cell_background(cell, "F1F5F9")
                cell.paragraphs[0].runs[0].bold = True
                cell.paragraphs[0].runs[0].font.size = Pt(9)
                cell.paragraphs[0].runs[0].font.color.rgb = RGBColor(0, 31, 63)
            else:
                set_cell_background(cell, "FFFFFF")
                cell.paragraphs[0].runs[0].font.size = Pt(8.5)
                cell.paragraphs[0].runs[0].font.color.rgb = RGBColor(51, 65, 85)
            set_cell_margins(cell, top=100, bottom=100, left=120, right=120)

    doc.add_page_break()

    # =========================================================================
    # INTRODUCCIÓN GENERAL E ÍNDICE
    # =========================================================================
    doc.add_heading("Introducción", level=1)
    doc.add_paragraph(
        "El aplicativo de Gestión y Emisión de Carnets Digitales (Carnet Digital) ha sido concebido para modernizar, "
        "agilizar y garantizar la trazabilidad en la identificación de colaboradores, estudiantes, contratistas y visitantes "
        "dentro de las organizaciones e instituciones, bajo estrictos estándares de seguridad digital y cumplimiento normativo."
    )
    doc.add_paragraph(
        "A través de esta herramienta tecnológica, las empresas suscritas (tenants) pueden administrar de forma autónoma su estructura "
        "organizacional (sedes, áreas y cargos), registrar su directorio de personal, diseñar plantillas personalizadas mediante un editor "
        "gráfico interactivo y emitir credenciales digitales seguras provistas de código QR criptográfico para control de acceso físico y lógico en tiempo real."
    )
    doc.add_paragraph(
        "Desarrollado en su capa de presentación con React 19 y empaquetado con Vite, el aplicativo se comunica mediante protocolos seguros REST "
        "con servicios de Node.js y bases de datos PostgreSQL. Esta arquitectura asegura tiempos de respuesta inferiores a 100 milisegundos, "
        "soporte para alta disponibilidad en cloud computing o entornos on-premises, y una experiencia fluida tanto en dispositivos de escritorio como móviles."
    )

    doc.add_heading("Contenido general del documento", level=2)
    tbl_index = doc.add_table(rows=1, cols=3)
    headers_index = ["Fase", "Capítulo / Módulo", "Público objetivo"]
    data_index = [
        ("Fase 1", "Capítulo 1: Arquitectura y topología del frontend", "Ingenieros de software y líderes técnicos"),
        ("Fase 1", "Capítulo 2: Seguridad, sesiones y contextos de autenticación", "Arquitectos de software y seguridad TIC"),
        ("Fase 1", "Capítulo 3: Catálogo maestro de rutas y matriz RBAC", "Desarrolladores frontend y analistas QA"),
        ("Fase 1", "Capítulo 4: Capa de servicios y consumo de API REST", "Desarrolladores full stack"),
        ("Fase 1", "Capítulo 5: Motor gráfico de credenciales (Canvas Engine)", "Ingenieros de interfaces y diseñadores UI"),
        ("Fase 2", "Módulo 1: Ingreso, registro y pasarela de pago (Figuras 1 a 9)", "Usuarios generales y directores comerciales"),
        ("Fase 2", "Módulo 2: Consola de administración de plataforma (Figuras 10 a 20)", "Superadministradores de sistema TIC"),
        ("Fase 2", "Módulo 3: Panel corporativo de empresa cliente (Figuras 21 a 35)", "Líderes de talento humano y operadores"),
        ("Anexo", "Preguntas frecuentes y resolución de problemas", "Mesa de ayuda y soporte técnico")
    ]
    format_table(tbl_index, [1.0, 3.8, 1.7], headers_index, data_index, "001F3F")

    doc.add_page_break()

    # =========================================================================
    # FASE 1: MANUAL TÉCNICO DE ARQUITECTURA FRONTEND
    # =========================================================================
    doc.add_heading("Fase 1. Manual técnico de arquitectura frontend", level=1)

    # 1.1 Stack
    doc.add_heading("1.1 Visión del sistema y catálogo tecnológico", level=2)
    doc.add_paragraph(
        "Carnet Digital es una aplicación de una sola página (Single Page Application o SPA) construida con React 19 y Vite. "
        "La interfaz proporciona tiempos de renderizado inmediatos, animaciones fluidas y un editor de plantillas en tiempo real "
        "con soporte para arquitecturas multi-inquilino (multi-tenant)."
    )

    tbl_tech = doc.add_table(rows=1, cols=3)
    headers_tech = ["Capa tecnológica", "Librería y versión", "Rol y responsabilidad en la plataforma"]
    data_tech = [
        ("Librería de interfaz", "React v19.x", "Renderizado reactivo de componentes, gestión declarativa de estados y ciclo de vida."),
        ("Herramienta de compilación", "Vite v5.x / Rollup", "Entorno de desarrollo rápido con Hot Module Replacement (HMR) y empaquetado optimizado."),
        ("Enrutamiento SPA", "React Router DOM v6.x / v7.x", "Navegación del lado del cliente, rutas protegidas por rol y layouts anidados."),
        ("Cliente HTTP", "Axios v1.x", "Interceptores para inyección automática de cabeceras Bearer JWT y tratamiento de errores 401/403."),
        ("Animaciones y transiciones", "Framer Motion v11.x", "Transición entre vistas, modales animados y microinteracciones de alta fidelidad."),
        ("Iconografía vectorial", "Lucide React", "Conjunto de íconos vectoriales SVG limpios y optimizados para reducir el tamaño final del paquete."),
        ("Motor gráfico nativo", "HTML5 Canvas API", "Renderizado de capas de carnet (imágenes, textos dinámicos y códigos QR) con exportación a 300 DPI.")
    ]
    format_table(tbl_tech, [1.6, 1.8, 3.1], headers_tech, data_tech)

    add_callout(doc, "El entorno de ejecución para el empaquetado de producción requiere Node.js 20 LTS o superior y el gestor de paquetes npm 10+.", "NOTE")

    # 1.2 Topología
    doc.add_heading("1.2 Estructura y topología modular del código fuente", level=2)
    doc.add_paragraph(
        "El código fuente en el directorio `src/` sigue el principio de separación de responsabilidades y modularidad orientada al dominio:"
    )

    tbl_dirs = doc.add_table(rows=1, cols=3)
    headers_dirs = ["Directorio", "Responsabilidad técnica", "Componentes destacados"]
    data_dirs = [
        ("src/core/", "Servicios compartidos, clientes HTTP y contexto global de sesión.", "AuthContext.jsx, apiClient.js, storage.js"),
        ("src/layouts/", "Estructuras de navegación y marcos de diseño visual.", "public/Layout.jsx, system/SystemLayout.jsx, company/CompanyLayout.jsx"),
        ("src/modules/public/", "Módulos de acceso libre sin autenticación previa.", "home/ (Landing page), blog/, auth/ (Login, Registro, Checkout)"),
        ("src/modules/system/", "Consola de administración global para superusuarios.", "home/, companies/, plans/, subscriptions/, users/, roles/, parameters/"),
        ("src/modules/company/", "Consola operativa de la empresa cliente (tenant).", "home/, headquarters/, areas/, roles/, users/, templates/, carnets/, parameter/, subscription/"),
        ("src/routes/", "Configuración centralizada de rutas y protección de accesos.", "PublicRoutes.jsx, SystemRoutes.jsx, CompanyRoutes.jsx, ProtectedRoute.jsx"),
        ("src/components/", "Elementos visuales genéricos y reutilizables.", "ui/ (Botones, Modales, Inputs), common/ (Pagination, LimitAlert)")
    ]
    format_table(tbl_dirs, [1.6, 2.4, 2.5], headers_dirs, data_dirs)

    # 1.3 Seguridad y JWT
    doc.add_heading("1.3 Gestión de seguridad, sesiones y matriz RBAC", level=2)
    doc.add_paragraph(
        "La autenticación utiliza tokens web JSON (JSON Web Tokens o JWT). Al iniciar sesión con credenciales válidas, "
        "el backend genera un token firmado que el cliente almacena en el navegador (`localStorage`) y transmite en cada petición HTTP "
        "mediante la cabecera `Authorization: Bearer <token>`."
    )
    doc.add_paragraph(
        "El contexto reactivo `AuthContext` decodifica la información del token y clasifica al usuario en dos alcances estrictos:\n"
        "• Tipo sistema (`type: 'system'`): Operador o superadministrador de la plataforma global. Su entrada natural es `/system/home`.\n"
        "• Tipo empresa (`type: 'company'`): Administrador corporativo vinculado a una organización (`company_id`). El sistema carga su logotipo y paleta de colores personalizada y lo redirige a `/company/home`."
    )

    add_callout(doc, "Si un usuario con rol de empresa intenta navegar manualmente hacia una ruta de `/system/*`, el componente `ProtectedRoute` deniega la acción y lo regresa a su panel autorizado.", "ALERT")

    # 1.4 Catálogo de rutas
    doc.add_heading("1.4 Catálogo de rutas del frontend y niveles de acceso", level=2)
    tbl_routes = doc.add_table(rows=1, cols=4)
    headers_routes = ["Ruta del frontend", "Nivel de acceso", "Estructura de layout", "Propósito en el sistema"]
    data_routes = [
        ("/", "Público", "Public Layout", "Página principal con catálogo interactivo y propuesta comercial."),
        ("/blog", "Público", "Public Layout", "Portal de artículos pedagógicos sobre identificación digital."),
        ("/login", "Público", "Standalone", "Formulario de inicio de sesión único para empresas y superusuarios."),
        ("/payment-success", "Público", "Standalone", "Confirmación de transacción aprobada y redirección."),
        ("/system/home", "Admin Sistema", "System Layout", "Panel principal con indicadores macro, cobros y empresas."),
        ("/system/companies", "Admin Sistema", "System Layout", "Directorio maestro de organizaciones suscritas."),
        ("/system/plans", "Admin Sistema", "System Layout", "Administración de tarifas comerciales y cuotas de carnets."),
        ("/system/subscriptions", "Admin Sistema", "System Layout", "Supervisión de transacciones y estados de cobranza."),
        ("/system/users", "Admin Sistema", "System Layout", "Gestión de operadores del núcleo de la plataforma."),
        ("/system/roles", "Admin Sistema", "System Layout", "Matriz de privilegios para usuarios del sistema."),
        ("/system/parameters", "Admin Sistema", "System Layout", "Control de banners promocionales y modo de mantenimiento."),
        ("/company/home", "Admin Empresa", "Company Layout", "Panel de bienvenida, estadísticas internas y consumo de cupos."),
        ("/company/headquarters", "Admin Empresa", "Company Layout", "Registro y administración de sedes físicas."),
        ("/company/areas", "Admin Empresa", "Company Layout", "Estructura departamental asociada a sedes."),
        ("/company/cargos", "Admin Empresa", "Company Layout", "Catálogo de puestos de trabajo organizacionales."),
        ("/company/users", "Admin Empresa", "Company Layout", "Directorio completo de colaboradores y personal."),
        ("/company/templates", "Admin Empresa", "Company Layout", "Galería de plantillas y acceso al editor Canvas."),
        ("/company/carnets", "Admin Empresa", "Company Layout", "Emisión masiva o individual, descarga en PDF y revocaciones."),
        ("/company/parameter", "Admin Empresa", "Company Layout", "Personalización de marca corporativa (white-labeling)."),
        ("/company/subscription", "Admin Empresa", "Company Layout", "Consulta del plan contratado y recarga de cuotas.")
    ]
    format_table(tbl_routes, [1.8, 1.2, 1.3, 2.2], headers_routes, data_routes)

    # 1.5 Servicios API
    doc.add_heading("1.5 Capa de servicios y comunicación REST", level=2)
    doc.add_paragraph(
        "Las peticiones hacia los microservicios backend se realizan mediante módulos de servicio especializados, "
        "garantizando desacoplamiento, trazabilidad y manejo uniforme de respuestas:"
    )

    tbl_services = doc.add_table(rows=1, cols=4)
    headers_services = ["Servicio frontend", "Módulo correspondiente", "Operaciones clave", "Puntos de enlace (Endpoints)"]
    data_services = [
        ("authService", "Autenticación", "login(), register(), verifyToken()", "POST /api/auth/login, POST /api/auth/register"),
        ("companyService", "Empresas", "getCompanies(), create(), update(), toggle()", "GET/POST/PUT /api/companies"),
        ("planService", "Planes del sistema", "getAllPlans(), createPlan(), updatePlan()", "GET/POST/PUT /api/plans"),
        ("subscriptionService", "Suscripciones", "getAllSubscriptions(), getStats()", "GET /api/subscriptions"),
        ("systemUserService", "Usuarios del sistema", "getAllUsers(), createUser(), deleteUser()", "GET/POST/DELETE /api/system-users"),
        ("systemRolService", "Roles del sistema", "getAllRoles(), createRole(), updateRole()", "GET/POST /api/system-roles"),
        ("headquarterService", "Sedes de empresa", "getAll(), create(), update(), delete()", "GET/POST/PUT /api/headquarters"),
        ("areaService", "Áreas de empresa", "getAll(), create(), update(), delete()", "GET/POST/PUT /api/areas"),
        ("cargoService", "Cargos laborales", "getAll(), create(), update(), delete()", "GET/POST/PUT /api/cargos"),
        ("colaboradorService", "Directorio de personal", "getAll(), create(), update(), delete()", "GET/POST/PUT /api/colaboradores"),
        ("templateService", "Plantillas de carnet", "getAllTemplates(), saveTemplate()", "GET/POST /api/templates"),
        ("carnetService", "Carnetización", "getAllCarnets(), emitir(), revocar()", "GET/POST/DELETE /api/carnets")
    ]
    format_table(tbl_services, [1.5, 1.5, 1.6, 1.9], headers_services, data_services)

    # 1.6 Motor Canvas
    doc.add_heading("1.6 Motor gráfico de credenciales (Canvas Engine)", level=2)
    doc.add_paragraph(
        "El diseñador de plantillas (`src/modules/company/templates/components/CanvasDesigner.jsx`) opera con un motor gráfico nativo "
        "sobre el elemento `<canvas>` de HTML5. Sus capacidades principales comprenden:\n"
        "• Sistema de capas ordenadas: Capa de fondo corporativo, logos vectoriales, fotografía del colaborador, textos y código QR.\n"
        "• Marcadores de posición dinámicos: Etiquetas como `{{nombres}}`, `{{apellidos}}`, `{{documento}}`, `{{cargo}}` y `{{sede}}` que el sistema reemplaza en tiempo real durante la emisión.\n"
        "• Generación de código QR seguro: Dibujado de códigos bidimensionales que contienen la dirección URL pública de verificación criptográfica.\n"
        "• Exportación de alta resolución: Salida gráfica rasterizada a 300 puntos por pulgada (DPI) adecuada para impresión física o visualización en teléfonos inteligentes."
    )

    doc.add_page_break()

    # =========================================================================
    # FASE 2: MANUAL DE USUARIO E INSTRUCTIVO PASO A PASO
    # =========================================================================
    doc.add_heading("Fase 2. Manual de usuario e instructivo paso a paso", level=1)
    doc.add_paragraph(
        "Esta fase describe el funcionamiento de cada una de las pantallas y ventanas modales de Carnet Digital. "
        "Sigue las instrucciones descritas según el numeral y la figura correspondiente."
    )

    # -------------------------------------------------------------------------
    # MÓDULO 1: PORTAL PÚBLICO Y ACCESO
    # -------------------------------------------------------------------------
    doc.add_heading("Módulo 1. Ingreso, registro y pasarela de pago", level=2)

    # Figura 1: Hero
    doc.add_heading("1. Portada principal y bienvenida comercial", level=3)
    doc.add_paragraph(
        "1. Abre tu navegador web favorito (Google Chrome, Microsoft Edge o Mozilla Firefox).\n"
        "2. Ingresa la dirección URL oficial de la plataforma: `http://localhost:5173/`.\n"
        "3. Observa la propuesta de valor y las opciones del menú de navegación superior."
    )
    add_screenshot(doc, "01_public_landing_hero.png", 1, "Pantalla de ingreso al portal web de Carnet Digital", max_width_inches=6.2)
    add_callout(doc, "Haz clic en el botón 'Empezar Ahora' para comunicarte vía WhatsApp con un asesor de servicio.", "TIP")

    # Figura 2: Servicios
    doc.add_heading("2. Catálogo interactivo de servicios y soluciones", level=3)
    doc.add_paragraph(
        "1. Desplázate hacia abajo en la página de inicio hasta la sección 'Servicios'.\n"
        "2. Explora las soluciones de credencialización disponibles: Identificación Digital, Control de Asistencia, Acceso a Sedes, Eventos y Certificaciones."
    )
    add_screenshot(doc, "02_public_landing_servicios.png", 2, "Catálogo de soluciones y funciones de credencialización", max_width_inches=5.8)

    # Figura 3: Planes
    doc.add_heading("3. Consulta de planes y tarifas comerciales", level=3)
    doc.add_paragraph(
        "1. Dirígete a la sección 'Suscripciones' mediante el menú superior.\n"
        "2. Alterna entre facturación mensual o anual con descuento especial.\n"
        "3. Elige el plan que mejor se adapte al volumen de colaboradores de tu organización."
    )
    add_screenshot(doc, "03_public_landing_planes.png", 3, "Selector de planes y tarifas de suscripción", max_width_inches=5.8)

    # Figura 4: Despliegue
    doc.add_heading("4. Opciones de arquitectura e infraestructura", level=3)
    doc.add_paragraph(
        "1. Revisa las modalidades de hospedaje técnico.\n"
        "2. Compara la solución Cloud SaaS gestionada frente al despliegue On-Premises en los servidores privados de tu institución."
    )
    add_screenshot(doc, "04_public_landing_despliegue.png", 4, "Comparativa de arquitectura Cloud SaaS frente a On-Premises", max_width_inches=5.8)

    # Figura 5: CTA y Footer
    doc.add_heading("5. Contacto comercial y pie de página", level=3)
    doc.add_paragraph(
        "1. Utiliza los canales directos del pie de página para resolver inquietudes técnicas o comerciales.\n"
        "2. Consulta las políticas de tratamiento de datos y términos de servicio institucional."
    )
    add_screenshot(doc, "05_public_landing_cta.png", 5, "Sección de llamado a la acción y enlaces institucionales", max_width_inches=6.0)

    # Figura 6: Blog
    doc.add_heading("6. Portal de novedades y artículos técnicos", level=3)
    doc.add_paragraph(
        "1. Haz clic en la opción 'Blog' de la barra superior.\n"
        "2. Filtra por categorías como 'Seguridad', 'Innovación' o 'Gestión' para consultar artículos de interés."
    )
    add_screenshot(doc, "06_public_blog.png", 6, "Directorio de publicaciones y artículos del blog institucional", max_width_inches=5.6)

    # Figura 7: Login
    doc.add_heading("7. Inicio de sesión en la plataforma", level=3)
    doc.add_paragraph(
        "1. Haz clic en el botón 'Entrar' ubicado en la esquina superior derecha.\n"
        "2. Escribe tu correo electrónico corporativo registrado.\n"
        "3. Escribe tu contraseña de acceso asignada.\n"
        "4. Haz clic en el botón **Iniciar Sesión**."
    )
    add_screenshot(doc, "07_auth_login.png", 7, "Ventana de autenticación e inicio de sesión seguro", max_width_inches=3.6)
    add_callout(doc, "Si la plataforma se encuentra en ventana de mantenimiento técnico, el sistema mostrará un aviso informativo antes de permitir el ingreso.", "NOTE")

    # Figura 8: Registro Paso 1
    doc.add_heading("8. Registro de una nueva empresa cliente", level=3)
    doc.add_paragraph(
        "1. En la tarjeta de acceso, haz clic en el enlace '¿No tienes cuenta? Regístrate aquí'.\n"
        "2. Diligencia la razón social completa de tu empresa y el número de identificación tributaria (NIT).\n"
        "3. Escribe los datos de contacto del administrador y define una contraseña segura."
    )
    add_screenshot(doc, "08_auth_registro_paso1.png", 8, "Formulario inicial para dar de alta una nueva organización", max_width_inches=3.6)

    # Figura 9: Pago Aprobado
    doc.add_heading("9. Confirmación de suscripción exitosa", level=3)
    doc.add_paragraph(
        "1. Al finalizar el pago en la pasarela, el sistema mostrará la pantalla de confirmación exitosa con tu identificador de transacción.\n"
        "2. Haz clic en **Ir al Login** para iniciar sesión y acceder inmediatamente al panel corporativo."
    )
    add_screenshot(doc, "09_auth_pago_aprobado.png", 9, "Comprobante oficial de pago aprobado y activación de cuenta", max_width_inches=4.2)

    doc.add_page_break()

    # -------------------------------------------------------------------------
    # MÓDULO 2: SUPERADMIN DE SISTEMA
    # -------------------------------------------------------------------------
    doc.add_heading("Módulo 2. Consola SuperAdmin de plataforma", level=2)
    doc.add_paragraph(
        "Este módulo está reservado exclusivamente para el equipo de ingenieros y administradores centrales de la Dirección TIC."
    )

    # Figura 10: Dashboard Sistema
    doc.add_heading("10. Tablero de control y métricas globales", level=3)
    doc.add_paragraph(
        "1. Revisa los indicadores clave de desempeño: Empresas activas, ingresos mensuales recurrentes (MRR) y total de carnets emitidos.\n"
        "2. Examina la gráfica de crecimiento intermensual y el registro de eventos recientes."
    )
    add_screenshot(doc, "10_system_dashboard.png", 10, "Tablero de control principal de la plataforma central", max_width_inches=5.8)

    # Figura 11: Empresas
    doc.add_heading("11. Directorio de empresas suscritas", level=3)
    doc.add_paragraph(
        "1. Haz clic en la opción 'Empresas' del menú lateral izquierdo.\n"
        "2. Consulta el listado de empresas con su plan activo, fecha de expiración y estado operativo.\n"
        "3. Utiliza la barra de búsqueda para ubicar rápidamente una empresa por su nombre o NIT."
    )
    add_screenshot(doc, "11_system_empresas.png", 11, "Directorio maestro y buscador de empresas suscritas", max_width_inches=5.8)

    # Figura 12: Modal Crear Empresa
    doc.add_heading("12. Registro manual de una nueva empresa", level=3)
    doc.add_paragraph(
        "1. Haz clic en el botón superior **+ Nueva Empresa**.\n"
        "2. Completa los campos: Razón social, NIT, correo del administrador y teléfono.\n"
        "3. Selecciona el plan comercial inicial contratado.\n"
        "4. Haz clic en **Guardar Empresa** para crear el inquilino (tenant) en la base de datos."
    )
    add_screenshot(doc, "12_system_modal_crear_empresa.png", 12, "Ventana modal para registrar manualmente una organización", max_width_inches=4.8)

    # Figura 13: Planes
    doc.add_heading("13. Catálogo de planes comerciales", level=3)
    doc.add_paragraph(
        "1. Haz clic en 'Planes' en el menú lateral.\n"
        "2. Visualiza las tarifas vigentes, el límite de colaboradores y el número máximo de plantillas autorizadas."
    )
    add_screenshot(doc, "13_system_planes.png", 13, "Directorio de planes y condiciones comerciales del sistema", max_width_inches=5.8)

    # Figura 14: Modal Crear Plan
    doc.add_heading("14. Creación y edición de condiciones de plan", level=3)
    doc.add_paragraph(
        "1. Haz clic en el botón **+ Nuevo Plan**.\n"
        "2. Define el nombre del plan (ejemplo: Enterprise Pro), precio mensual y precio anual.\n"
        "3. Especifica el límite estricto de colaboradores y de plantillas.\n"
        "4. Haz clic en **Crear Plan**."
    )
    add_screenshot(doc, "13b_system_modal_crear_plan.png", 14, "Modal para parametrizar costos y cuotas de un plan", max_width_inches=4.4)

    # Figura 15: Suscripciones
    doc.add_heading("15. Auditoría de suscripciones y transacciones", level=3)
    doc.add_paragraph(
        "1. Ingresa a la sección 'Suscripciones'.\n"
        "2. Supervisa las renovaciones automáticas, estados de cuenta (al día, vencido) y facturas generadas."
    )
    add_screenshot(doc, "14_system_suscripciones.png", 15, "Monitor de facturación y transacciones de suscripción", max_width_inches=5.8)

    # Figura 16: Usuarios Sistema
    doc.add_heading("16. Administración de operadores del sistema", level=3)
    doc.add_paragraph(
        "1. Haz clic en 'Usuarios' dentro de la consola del sistema.\n"
        "2. Observa la nómina de operadores con privilegios de superadministrador."
    )
    add_screenshot(doc, "15_system_usuarios.png", 16, "Listado de operadores y técnicos del núcleo central", max_width_inches=5.8)

    # Figura 17: Modal Crear Usuario
    doc.add_heading("17. Alta de operador del sistema", level=3)
    doc.add_paragraph(
        "1. Haz clic en **+ Nuevo Usuario**.\n"
        "2. Escribe los nombres, correo institucional y contraseña temporal.\n"
        "3. Asigna el rol correspondiente (ejemplo: Administrador o Soporte).\n"
        "4. Haz clic en **Guardar Usuario**."
    )
    add_screenshot(doc, "15b_system_modal_crear_usuario.png", 17, "Ventana para crear cuentas de operadores técnicos", max_width_inches=4.2)

    # Figura 18: Roles Sistema
    doc.add_heading("18. Matriz de roles y privilegios", level=3)
    doc.add_paragraph(
        "1. Ingresa a la opción 'Roles'.\n"
        "2. Consulta las definiciones de perfiles del sistema y su alcance de gestión."
    )
    add_screenshot(doc, "16_system_roles.png", 18, "Matriz de privilegios y perfiles de usuario del sistema", max_width_inches=5.8)

    # Figura 19: Modal Crear Rol
    doc.add_heading("19. Creación de un nuevo rol de sistema", level=3)
    doc.add_paragraph(
        "1. Haz clic en **+ Nuevo Rol**.\n"
        "2. Ingresa el identificador único y una descripción de responsabilidades.\n"
        "3. Guarda la configuración para aplicarla de inmediato a los operadores."
    )
    add_screenshot(doc, "16b_system_modal_crear_rol.png", 19, "Formulario modal para registrar un nuevo rol administrativo", max_width_inches=4.2)

    # Figura 20: Parámetros y Banners
    doc.add_heading("20. Parámetros globales y banners institucionales", level=3)
    doc.add_paragraph(
        "1. Haz clic en 'Parámetros' en el menú lateral.\n"
        "2. Activa o desactiva el switch de 'Permitir Inicio de Sesión' para labores de mantenimiento.\n"
        "3. Carga nuevas imágenes de banner y verifica su previsualización en tiempo real."
    )
    add_screenshot(doc, "17_system_parametros_banners.png", 20, "Panel de control de modo mantenimiento y banners informativos", max_width_inches=5.8)

    doc.add_page_break()

    # -------------------------------------------------------------------------
    # MÓDULO 3: PANEL DE EMPRESA (CLIENTE TENANT)
    # -------------------------------------------------------------------------
    doc.add_heading("Módulo 3. Panel corporativo de la empresa cliente", level=2)
    doc.add_paragraph(
        "Este módulo está destinado a los líderes de Talento Humano, Seguridad y Administradores de cada organización suscrita."
    )

    # Figura 21: Dashboard Empresa
    doc.add_heading("21. Tablero principal de la empresa", level=3)
    doc.add_paragraph(
        "1. Accede con tus credenciales de empresa.\n"
        "2. Observa el saludo institucional, el resumen de colaboradores activos y la distribución por sedes.\n"
        "3. Verifica la barra de progreso de cupos de carnets disponibles."
    )
    add_screenshot(doc, "18_company_dashboard.png", 21, "Tablero de control corporativo de la organización", max_width_inches=5.8)

    # Figura 22: Sedes
    doc.add_heading("22. Gestión de sedes y sucursales físicas", level=3)
    doc.add_paragraph(
        "1. Haz clic en la opción 'Sedes' del menú corporativo.\n"
        "2. Consulta las instalaciones físicas registradas (oficinas principales, plantas o sucursales)."
    )
    add_screenshot(doc, "19_company_sedes.png", 22, "Directorio maestro de sedes físicas registradas", max_width_inches=5.8)

    # Figura 23: Modal Crear Sede
    doc.add_heading("23. Registro de una nueva sede", level=3)
    doc.add_paragraph(
        "1. Haz clic en el botón **+ Nueva Sede**.\n"
        "2. Escribe el nombre de la sede (ejemplo: Sede Principal Calle 100), ciudad y dirección exacta.\n"
        "3. Haz clic en **Guardar Sede**."
    )
    add_screenshot(doc, "20_company_modal_crear_sede.png", 23, "Modal de ingreso de datos de una sede física", max_width_inches=4.2)

    # Figura 24: Áreas
    doc.add_heading("24. Estructura de áreas y departamentos", level=3)
    doc.add_paragraph(
        "1. Dirígete a la opción 'Áreas'.\n"
        "2. Examina los departamentos funcionales (Tecnología, Operaciones, Ventas) adscritos a cada sede."
    )
    add_screenshot(doc, "21_company_areas.png", 24, "Estructura departamental y áreas de trabajo de la empresa", max_width_inches=5.8)

    # Figura 25: Modal Crear Área
    doc.add_heading("25. Creación de un área o departamento", level=3)
    doc.add_paragraph(
        "1. Haz clic en **+ Nueva Área**.\n"
        "2. Escribe la denominación del área.\n"
        "3. Selecciona en el menú desplegable la sede física a la cual pertenece.\n"
        "4. Haz clic en **Guardar Área**."
    )
    add_screenshot(doc, "22_company_modal_crear_area.png", 25, "Formulario modal para crear y asociar un área", max_width_inches=4.2)

    # Figura 26: Cargos
    doc.add_heading("26. Catálogo de cargos y puestos de trabajo", level=3)
    doc.add_paragraph(
        "1. Haz clic en 'Cargos' en el menú lateral.\n"
        "2. Consulta el listado de cargos laborales organizados por área funcional."
    )
    add_screenshot(doc, "23_company_cargos.png", 26, "Catálogo de cargos laborales activos en la organización", max_width_inches=5.8)

    # Figura 27: Modal Crear Cargo
    doc.add_heading("27. Alta de un nuevo cargo", level=3)
    doc.add_paragraph(
        "1. Haz clic en **+ Nuevo Cargo**.\n"
        "2. Ingresa el título del puesto (ejemplo: Ingeniero de Software Senior).\n"
        "3. Selecciona el área a la que pertenece.\n"
        "4. Haz clic en **Guardar Cargo**."
    )
    add_screenshot(doc, "24_company_modal_crear_cargo.png", 27, "Ventana modal para registrar un cargo laboral", max_width_inches=4.2)

    # Figura 28: Colaboradores
    doc.add_heading("28. Directorio maestro de colaboradores", level=3)
    doc.add_paragraph(
        "1. Haz clic en la opción 'Colaboradores'.\n"
        "2. Visualiza a todo el personal con su fotografía, documento de identidad, cargo, sede y estado de carnetización.\n"
        "3. Utiliza los filtros superiores para segmentar por área o cargo."
    )
    add_screenshot(doc, "25_company_colaboradores.png", 28, "Directorio general de colaboradores con opciones de búsqueda", max_width_inches=5.8)

    # Figura 29: Modal Crear Colaborador
    doc.add_heading("29. Registro completo de un colaborador", level=3)
    doc.add_paragraph(
        "1. Haz clic en el botón superior **+ Nuevo Colaborador**.\n"
        "2. Selecciona el tipo de documento (Cédula de Ciudadanía, Pasaporte o Cédula de Extranjería) e ingresa el número.\n"
        "3. Escribe los nombres, apellidos, correo corporativo y número de celular.\n"
        "4. Asigna la sede, área y cargo correspondientes.\n"
        "5. Carga la fotografía formal del colaborador en formato JPG o PNG.\n"
        "6. Haz clic en **Guardar Colaborador**."
    )
    add_screenshot(doc, "26_company_modal_colaborador.png", 29, "Formulario modal integral para el registro de colaboradores", max_width_inches=5.0)

    # Figura 30: Plantillas
    doc.add_heading("30. Galería de plantillas de credenciales", level=3)
    doc.add_paragraph(
        "1. Haz clic en 'Plantillas' en el menú corporativo.\n"
        "2. Revisa los diseños creados para personal de planta, contratistas, directivos o eventos especiales."
    )
    add_screenshot(doc, "27_company_plantillas.png", 30, "Galería de plantillas visuales de carnetización", max_width_inches=5.8)

    # Figura 31: Diseñador Canvas
    doc.add_heading("31. Diseñador interactivo de plantillas de carnet", level=3)
    doc.add_paragraph(
        "1. Haz clic en **+ Nueva Plantilla** o selecciona una existente para editarla.\n"
        "2. Sube la imagen de fondo institucional de tu organización.\n"
        "3. Arrastra y ubica las capas de texto dinámico: `{{nombres}}`, `{{apellidos}}`, `{{cargo}}` y `{{sede}}`.\n"
        "4. Ajusta la posición y tamaño del recuadro para el código QR y la fotografía.\n"
        "5. Haz clic en **Guardar Plantilla** para habilitarla en las emisiones."
    )
    add_screenshot(doc, "28_company_disenador_canvas.png", 31, "Editor gráfico interactivo Canvas para diseñar credenciales", max_width_inches=5.5)

    # Figura 32: Carnetización
    doc.add_heading("32. Control de emisión y descarga de carnets", level=3)
    doc.add_paragraph(
        "1. Haz clic en la opción 'Carnetización'.\n"
        "2. Consulta los carnets emitidos, vigencias y fechas de expedición.\n"
        "3. Haz clic en el ícono de descarga para obtener la credencial en PDF de alta resolución.\n"
        "4. En caso de retiro del funcionario o pérdida del carnet, utiliza la opción **Revocar**."
    )
    add_screenshot(doc, "29_company_carnetizacion.png", 32, "Listado de carnets emitidos con opciones de descarga y revocación", max_width_inches=5.8)
    add_callout(doc, "La revocación de una credencial invalida de inmediato el código QR ante cualquier lector de control de acceso físico.", "WARNING")

    # Figura 33: Modal Emitir Carnet
    doc.add_heading("33. Emisión individual de credencial", level=3)
    doc.add_paragraph(
        "1. Haz clic en el botón **+ Emitir Carnet**.\n"
        "2. Elige el colaborador en el selector desplegable.\n"
        "3. Elige la plantilla visual aprobada.\n"
        "4. Fija la fecha de expiración del carnet.\n"
        "5. Haz clic en **Emitir Credencial**."
    )
    add_screenshot(doc, "29b_company_modal_emitir_carnet.png", 33, "Modal de emisión individual y generación de código QR", max_width_inches=4.4)
    add_callout(doc, "Si la cuota de carnets contratados se encuentra agotada, el sistema bloqueará la emisión y te solicitará contactar a tu asesor para recargar cupos.", "ALERT")

    # Figura 34: Branding
    doc.add_heading("34. Configuración de marca corporativa (White-Labeling)", level=3)
    doc.add_paragraph(
        "1. Ingresa a la opción 'Configuración de Marca'.\n"
        "2. Personaliza el color primario y secundario con el selector de color o código hexadecimal.\n"
        "3. Elige la tipografía institucional preferida.\n"
        "4. Carga el archivo del logotipo oficial de tu empresa.\n"
        "5. Haz clic en **Guardar Cambios** para aplicar la identidad visual en todo tu entorno."
    )
    add_screenshot(doc, "30_company_branding_parametros.png", 34, "Módulo de configuración de marca, colores y logotipo de empresa", max_width_inches=5.8)

    # Figura 35: Suscripción Empresa
    doc.add_heading("35. Consulta de plan, cuotas y recarga de cupos", level=3)
    doc.add_paragraph(
        "1. Haz clic en la opción 'Suscripción'.\n"
        "2. Verifica el plan vigente, costo y fecha de renovación programada.\n"
        "3. Consulta la barra de consumo de carnets emitidos sobre el límite total permitido.\n"
        "4. Haz clic en **Actualizar Plan / Recargar Cupos** si requieres emitir credenciales adicionales."
    )
    add_screenshot(doc, "31_company_suscripcion.png", 35, "Panel de control del plan contratado y consumo de cupos", max_width_inches=5.8)

    # =========================================================================
    # ANEXO: PREGUNTAS FRECUENTES Y RESOLUCIÓN DE PROBLEMAS
    # =========================================================================
    doc.add_page_break()
    doc.add_heading("Anexo. Preguntas frecuentes y solución de problemas", level=1)
    doc.add_paragraph(
        "A continuación se presentan las situaciones más habituales y las acciones correctivas recomendadas:"
    )

    tbl_faq = doc.add_table(rows=1, cols=3)
    headers_faq = ["Situación o duda frecuente", "Causa probable", "Solución recomendada paso a paso"]
    data_faq = [
        (
            "El botón 'Emitir Carnet' se encuentra inhabilitado",
            "La empresa ha alcanzado el límite máximo de carnets contratados en su plan.",
            "Ingresa al módulo 'Suscripción' y selecciona la opción de recarga de cupos o cambio a un plan superior."
        ),
        (
            "Al escanear el carnet con el teléfono, se indica credencial inválida",
            "El carnet ha superado su fecha de expiración o fue revocado por el administrador.",
            "Verifica el estado en el módulo 'Carnetización'. Si está revocado, emite una nueva credencial con fecha vigente."
        ),
        (
            "No se visualizan las sedes al crear un área",
            "No existen sedes creadas previamente en la organización.",
            "Ingresa primero al módulo 'Sedes', registra al menos una sucursal física y posteriormente asocia el área."
        ),
        (
            "El logotipo no se visualiza correctamente en el carnet",
            "El formato o las dimensiones de la imagen cargada no son transparentes o superan el tamaño sugerido.",
            "Sube el logotipo en formato PNG con fondo transparente y resolución de 500x500 píxeles desde el módulo 'Configuración de Marca'."
        )
    ]
    format_table(tbl_faq, [1.8, 2.0, 2.7], headers_faq, data_faq, "001F3F")

    # Guardar documento
    print(f"Guardando documento Word oficial en: {DOCX_OUTPUT_PATH} ...")
    doc.save(DOCX_OUTPUT_PATH)
    print("¡DOCUMENTO GENERADO EXITOSAMENTE CON ESTILO MICROSOFT Y PLANTILLA INSTITUCIONAL!")

if __name__ == "__main__":
    main()
